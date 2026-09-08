"""Independent consumer validation of the actual synthetic browser XLSX export."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import math
import warnings
import openpyxl

path = Path('test-output/history-export-synthetic.xlsx')
with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter('always')
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
assert not caught, [str(w.message) for w in caught]
assert formulas.sheetnames == ['SUMINISTROS', 'CUPS 1', 'RESUMEN EMPRESA', 'PERIODOS', 'DETALLE P1-P6']
for name in ['CUPS 1', 'RESUMEN EMPRESA']:
    ws = formulas[name]
    assert len(ws._images) == 3, (name, len(ws._images))
    anchors = [image.anchor for image in ws._images]
    assert all(type(a).__name__ == 'TwoCellAnchor' for a in anchors)
    assert all(a.to.row - a._from.row == 15 for a in anchors)
    assert all(anchors[i+1]._from.row > anchors[i].to.row for i in range(2))
    assert anchors[0]._from.row >= 20
    assert values[name]['B18'].value == 450
    assert values[name]['C18'].value == 120
    assert math.isclose(values[name]['D18'].value, .1)
    assert math.isclose(values[name]['E18'].value, 120/450)
    assert values[name]['B7'].value is None
    assert values[name]['D7'].value in [None, '']
    assert ws['B5'].value.startswith('=SUMIFS(')
    assert ws['A5'].value == 'ENERO 2026'
    assert ws['A17'].value == 'ENERO 2027'
    assert ws.sheet_view.showGridLines is False
    assert ws.freeze_panes == 'A5'
assert formulas['SUMINISTROS'].max_column == 10
assert formulas['PERIODOS']['A5'].comment is not None
assert 'A-JAN' in formulas['PERIODOS']['A5'].comment.text
for ws in values:
    for row in ws:
        for cell in row:
            assert cell.data_type != 'e', (ws.title, cell.coordinate, cell.value)
with ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
    ns = {'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}
    drawings = [n for n in z.namelist() if n.startswith('xl/drawings/drawing') and n.endswith('.xml')]
    assert len(drawings) == 2
    for name in drawings:
        root = ET.fromstring(z.read(name))
        assert len(root.findall('xdr:twoCellAnchor', ns)) == 3
        assert len(root.findall('xdr:oneCellAnchor', ns)) == 0
print('PASS: independent openpyxl load without warnings; six retained images; valid drawing anchors without overlaps; cached totals, unit prices, source notes and XML. Synthetic test only.')
